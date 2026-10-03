import os
import sys
import time
import json
import re
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT_DIR))
OUTPUT_DIR = ROOT_DIR / "dist" / "leads_output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

class LeadCaptureEngine:
    """GhostDesk 小红书/抖音/视频号评论区截流与意向客户挖掘引擎"""

    INTENT_KEYWORDS = {
        "询价/报价": ["多少钱", "价格", "报价", "底价", "优惠", "怎么卖", "费用"],
        "求购/购买": ["求链接", "怎么买", "哪里买", "有链接吗", "下单", "求购", "求带", "买一套"],
        "招商/合作": ["代理", "批发", "合作", "采购", "贴牌", "批量", "定制", "怎么加盟"],
        "私聊/沟通": ["私我", "私信", "加我", "怎么联系", "聊聊", "微我", "留个联系方式"]
    }

    @classmethod
    def analyze_comment(cls, comment_text: str) -> dict:
        """分析单条评论是否具有高购买意向，提取电话、微信号与意向标签"""
        matched_tags = []
        matched_keywords = []
        for tag, kws in cls.INTENT_KEYWORDS.items():
            hit = [k for k in kws if k in comment_text]
            if hit:
                matched_tags.append(tag)
                matched_keywords.extend(hit)

        is_intent = len(matched_tags) > 0
        
        # 1. 提取手机号 (支持连字符和空格)
        phone_match = re.search(r"1[3-9]\d[\s-]?\d{4}[\s-]?\d{4}", comment_text)
        phone = re.sub(r"[\s-]", "", phone_match.group(0)) if phone_match else ""

        # 2. 提取微信号
        wechat = ""
        wx_match = re.search(r"(?:vx|wx|微信|加v|v信|V信|V|v)[:：\s]*([a-zA-Z0-9_-]{6,20})", comment_text, re.IGNORECASE)
        if wx_match:
            wechat = wx_match.group(1)

        # 综合联系方式
        contacts = []
        if phone:
            contacts.append(f"手机:{phone}")
        if wechat:
            contacts.append(f"微信:{wechat}")
        contact_str = " / ".join(contacts)

        # 如果直接留了联系方式，哪怕没命中关键词也是绝对的高意向！
        if contact_str:
            is_intent = True
            if "私聊/沟通" not in matched_tags:
                matched_tags.append("直接留资")

        return {
            "is_high_intent": is_intent,
            "intent_tags": matched_tags,
            "matched_keywords": matched_keywords,
            "detected_contact": contact_str,
            "sentiment": "high_lead" if is_intent else "normal"
        }


    @classmethod
    def parse_raw_text(cls, raw_text: str) -> list:
        """从杂乱的网页复制文本或多行文本流中智能解析提取用户与评论内容"""
        lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
        comments = []
        for line in lines:
            # 常见格式 1: "昵称：评论内容" 或 "昵称: 评论内容"
            if "：" in line or ":" in line:
                parts = re.split(r"[：:]", line, maxsplit=1)
                user = parts[0].strip()
                text = parts[1].strip()
                if len(user) < 30 and len(text) > 0:
                    comments.append({"user": user, "text": text})
                    continue
            # 常见格式 2: 纯文本行
            if len(line) >= 2:
                comments.append({"user": "意向访客", "text": line})
        return comments

    @classmethod
    def capture_from_clipboard(cls, source_platform: str = "小红书") -> dict:
        """一键从 Windows 剪贴板读取复制的评论内容并自动执行截流挖掘"""
        import pyperclip
        text = pyperclip.paste()
        if not text or not text.strip():
            return {
                "status": "warning",
                "count": 0,
                "message": "剪贴板为空，请先在小红书/抖音等网页中复制评论文字！"
            }
        comments = cls.parse_raw_text(text)
        return cls.process_leads_batch(source_platform, comments)

    @classmethod
    def open_platform_search(cls, platform: str, keyword: str) -> dict:
        """自动唤醒浏览器定位到小红书或抖音的目标搜索/笔记页面"""
        import webbrowser
        import urllib.parse
        encoded = urllib.parse.quote(keyword)
        if "抖音" in platform:
            url = f"https://www.douyin.com/search/{encoded}"
        elif "快手" in platform:
            url = f"https://www.kuaishou.com/search/video?searchKey={encoded}"
        else:
            # 默认小红书搜索
            url = f"https://www.xiaohongshu.com/search_result?keyword={encoded}&source=web_search_result_notes"
        
        webbrowser.open(url)
        print(f"🌐 [截流特工] 已自动打开【{platform}】搜索页面: {url}", flush=True)
        return {
            "status": "success",
            "url": url,
            "message": f"已自动打开 {platform} 并定位搜索「{keyword}」"
        }

    @classmethod
    def scrape_url_live(cls, url: str, source_platform: str = "小红书") -> dict:
        """
        利用开源神器 DrissionPage 真实抓取网页内容与评论流
        支持自动清洗 HTML，提取正文与高意向线索
        """
        print(f"\n🌐 [DrissionPage 引擎] 启动真实网页抓取: {url}", flush=True)
        try:
            from DrissionPage import SessionPage
            from lxml import html
            page = SessionPage()
            page.get(url, timeout=10)
            
            # 使用 lxml 极速深度提取正文与评论
            tree = html.fromstring(page.html)
            raw_text = tree.text_content()
            lines = [l.strip() for l in raw_text.splitlines() if len(l.strip()) > 3]
            print(f"📄 [DrissionPage 引擎] 网页抓取成功，提取到 {len(lines)} 行有效文本！", flush=True)
            comments = cls.parse_raw_text("\n".join(lines[:100]))
            return cls.process_leads_batch(source_platform, comments)
        except Exception as e:
            print(f"⚠️ [DrissionPage 引擎] 真实抓取遇阻: {e}", flush=True)
            return {
                "status": "error",
                "message": f"真实网页抓取异常: {e}"
            }

    @classmethod
    def process_leads_batch(cls, source_platform: str, comments_data: list) -> dict:

        """批量处理截流抓取到的评论列表，并联动'咔哒'生成定制话术"""
        from skills.kada_reply.engine import KadaReplyEngine
        from skills.feishu_sync.engine import FeishuSyncEngine

        leads = []
        for item in comments_data:
            user = item.get("user", "匿名用户")
            text = item.get("text", "")
            analysis = cls.analyze_comment(text)
            if analysis["is_high_intent"]:
                # 联动咔哒生成高情商回复话术
                reply_pitch = KadaReplyEngine.generate_reply_text(user, text)
                leads.append({
                    "platform": source_platform,
                    "user": user,
                    "comment": text,
                    "intent_tags": analysis["intent_tags"],
                    "contact": analysis["detected_contact"],
                    "keywords": analysis["matched_keywords"],
                    "recommended_reply": reply_pitch,
                    "captured_at": time.strftime("%Y-%m-%d %H:%M:%S")
                })

        # 保存至本地截流客户库 (JSON)
        timestamp = time.strftime('%Y%m%d_%H%M%S')
        out_file = OUTPUT_DIR / f"leads_{timestamp}.json"
        with open(out_file, "w", encoding="utf-8") as f:
            json.dump(leads, f, indent=2, ensure_ascii=False)

        # 导出为专业 Excel 表格 (.xlsx)
        excel_file = OUTPUT_DIR / f"高意向客户线索表_{timestamp}.xlsx"
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "截流高意向客户"

            headers = ["序号", "来源平台", "客户昵称", "意向分类", "联系方式(手机/微信)", "评论原文", "咔哒推荐回复话术", "捕获时间"]
            ws.append(headers)

            # 样式美化：企业蓝表头 + 居中加粗
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for i, lead in enumerate(leads, 1):
                ws.append([
                    i,
                    lead.get("platform", ""),
                    lead.get("user", ""),
                    ", ".join(lead.get("intent_tags", [])),
                    lead.get("contact", "") or "未留(私信转化)",
                    lead.get("comment", ""),
                    lead.get("recommended_reply", ""),
                    lead.get("captured_at", "")
                ])

            # 自适应列宽
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len * 1.5, 12), 42)

            wb.save(str(excel_file))
            print(f"📊 [截流特工] 已导出 Excel 报表: {excel_file.name}")
        except Exception as e:
            print(f"⚠️ [截流特工] Excel 导出跳过: {e}")

        # 联动飞书记录
        if leads:
            FeishuSyncEngine.append_bitable_record({
                "record_type": "高意向客户截流",
                "platform": source_platform,
                "leads_count": len(leads),
                "leads_preview": [f"{l['user']}: {l['comment'][:20]}" for l in leads[:3]],
                "excel_file": str(excel_file.name),
                "saved_file": str(out_file.name)
            })

        print(f"🎯 [截流特工] 在【{source_platform}】共抓取并筛选出 {len(leads)} 个高意向客户！已保存至: {out_file.name}")
        return {
            "status": "success",
            "count": len(leads),
            "file": str(out_file.resolve()),
            "excel_file": str(excel_file.resolve()) if excel_file.exists() else "",
            "leads": leads
        }

if __name__ == "__main__":
    test_comments = [
        {"user": "数码小达人", "text": "这个对讲机多少钱啊？求链接！"},
        {"user": "路过的小猫", "text": "拍得挺好看的。"},
        {"user": "电商老张", "text": "能批量管理多台电脑吗？怎么联系你？私我"},
        {"user": "李先生", "text": "我的微信 13812345678 加我详聊合作"}
    ]
    res = LeadCaptureEngine.process_leads_batch("小红书", test_comments)
    print("测试结果:\n", json.dumps(res, indent=2, ensure_ascii=False))
